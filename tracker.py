# -*- coding: utf-8 -*-
"""게임 유저 반응 트래커 (게임별 값은 config.json).

  python tracker.py steam       # SteamDB Rating 기록 (15분마다)
  python tracker.py community   # Reddit + DC + Steam 토론 30분 창 분석 + 검토용 엑셀 (30분마다)
  python tracker.py build       # 대시보드(index.html)만 재생성
  python tracker.py learn       # 라벨링 엑셀 다시 읽고 학습 상태 출력
  python tracker.py rescore     # 현재 모델로 과거 수집분 전체 재분류

시각은 모두 KST. 30분 창은 정시/30분 경계에 맞춘다(14:31 실행 → 14:00~14:30 분석).
"""
import html
import json
import math
import os
import random
import re
import sys
import threading
import time
import traceback
from datetime import datetime, timedelta, timezone

import requests
from bs4 import BeautifulSoup

import sentiment as S
import translate as TR

KST = timezone(timedelta(hours=9))
BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, 'data')
ITEMS = os.path.join(DATA, 'items')
LABEL_DIR = os.path.join(BASE, '라벨링')
# GitHub Actions 에서 돌 때(TRACKER_CLOUD=1): data/ 는 data 브랜치를 체크아웃한 폴더라 로그도 그 안에 둬야 남는다.
CLOUD = os.environ.get('TRACKER_CLOUD') == '1'
LOG_DIR = os.path.join(DATA, 'logs') if CLOUD else os.path.join(BASE, 'logs')
LABELS_ENC = os.path.join(BASE, 'labels.enc')  # 클라우드용 라벨(판단 이유 포함) — 공개 저장소라 암호화
LABELS_KEY_FILE = os.path.join(os.path.expanduser('~'), '.config', 'github', 'labels_key.txt')
for d in (DATA, ITEMS, LABEL_DIR, LOG_DIR):
    os.makedirs(d, exist_ok=True)


def item_path(slot_end):
    """수집 항목은 30분 창마다 한 파일(data/items/YYYY-MM-DD/HHMM.jsonl, HHMM=창 끝).
    월 단위 한 파일은 하루 ~12MB씩 커져 GitHub 파일 한도(100MB)를 넘기 때문(2026-10-08)."""
    if isinstance(slot_end, str):
        slot_end = datetime.fromisoformat(slot_end)
    return os.path.join(ITEMS, f'{slot_end:%Y-%m-%d}', f'{slot_end:%H%M}.jsonl')


def item_files():
    out = []
    for root, _, files in os.walk(ITEMS):
        out += [os.path.join(root, f) for f in files if f.endswith('.jsonl')]
    return sorted(out)

# 게임별 설정. 소스를 안 쓰려면 해당 값을 null/빈값으로 (예: 서브레딧이 없으면 reddit_sub: null)
CFG = json.load(open(os.path.join(BASE, 'config.json'), encoding='utf-8'))
GAME = CFG.get('game', '게임')
APP_ID = CFG.get('steam_appid')
SUB = CFG.get('reddit_sub')
DC_ID = CFG.get('dc_id')
DC_TYPE = CFG.get('dc_type', 'mgallery')  # mgallery(마이너) / board(정식) / mini(미니)
S.extend_topics(CFG.get('topic_extra'))
WINDOW = timedelta(minutes=30)
REVIEW_N = 30
REVIEW_EVERY = int(CFG.get('review_every_min', 60))  # 라벨링 엑셀 주기(분, 30의 배수, 0=중단). 수집 주기와 별개
# 엑셀 '이유 유형' 드롭다운. 자유 기술은 '이유 설명' 칸에.
REASON_TYPES = ['과금/BM', '버그/서버/최적화', '밸런스/직업', '작업장/핵/봇', '운영/소통', '콘텐츠 부족/숙제',
                '비꼼/반어', '타 게임 비교', '욕설이지만 게임 비판', '게임과 무관', '기대/호평', '기타']
UA = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
      '(KHTML, like Gecko) Chrome/130.0 Safari/537.36')
BOT_AUTHORS = {a.lower() for a in CFG.get('reddit_bots', ['automoderator'])}
SOURCES = tuple(s for s, on in (('reddit', SUB), ('dc', DC_ID), ('steamforum', APP_ID and CFG.get('steam_forums') is not False)) if on)
SRC_NAME = {'reddit': 'Reddit', 'dc': 'DC', 'steamforum': 'Steam 토론'}

_log_lock = threading.Lock()


def log(msg):
    line = f'[{datetime.now(KST):%Y-%m-%d %H:%M:%S}] {msg}'
    with _log_lock:
        try:
            print(line)
        except UnicodeEncodeError:
            print(line.encode('ascii', 'replace').decode())
        with open(os.path.join(LOG_DIR, f'{datetime.now(KST):%Y-%m}.log'), 'a', encoding='utf-8') as f:
            f.write(line + '\n')


def append_jsonl(path, rec):
    with open(path, 'a', encoding='utf-8') as f:
        f.write(json.dumps(rec, ensure_ascii=False) + '\n')


def read_jsonl(path):
    if not os.path.exists(path):
        return []
    out = []
    with open(path, encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    out.append(json.loads(line))
                except json.JSONDecodeError:
                    pass
    return out


def iso(dt):
    return dt.astimezone(KST).strftime('%Y-%m-%dT%H:%M:%S')


# ====================================================================== Steam
def steamdb_rating(pos, total):
    """SteamDB 공식 산식 (steamdb.info/blog/steamdb-rating).
    rating = score - (score - 0.5) * 2^(-log10(total + 1))"""
    if total <= 0:
        return None
    score = pos / total
    return score - (score - 0.5) * 2 ** (-math.log10(total + 1))


def run_steam():
    """SteamDB 페이지는 Cloudflare(403)로 막혀 있어, 같은 원천인 Steam 리뷰 API 의
    전체 언어·전체 구매경로 긍/부정 수에 SteamDB 산식을 적용한다."""
    url = (f'https://store.steampowered.com/appreviews/{APP_ID}'
           '?json=1&language=all&purchase_type=all&num_per_page=0&filter=all')
    last_err = None
    for attempt in range(3):
        try:
            r = requests.get(url, headers={'User-Agent': UA}, timeout=20)
            q = r.json()['query_summary']
            break
        except Exception as e:  # noqa
            last_err = e
            time.sleep(5 * (attempt + 1))
    else:
        log(f'[steam] 실패: {last_err}')
        return 1
    pos, neg = q['total_positive'], q['total_negative']
    total = pos + neg
    path = os.path.join(DATA, 'steam.jsonl')
    prev = read_jsonl(path)
    prev = prev[-1] if prev else None
    rec = {
        'ts': iso(datetime.now(KST)),
        'positive': pos, 'negative': neg, 'total': total,
        'steam_pct': round(pos / total * 100, 2) if total else None,
        'steamdb_rating': round(steamdb_rating(pos, total) * 100, 2) if total else None,
        'desc': q.get('review_score_desc'),
    }
    if prev:
        rec['new_pos'] = pos - prev['positive']
        rec['new_neg'] = neg - prev['negative']
    append_jsonl(path, rec)
    log(f"[steam] SteamDB {rec['steamdb_rating']}% / Steam {rec['steam_pct']}% "
        f"({pos:,}+ / {neg:,}-) {rec['desc']}")
    return 0


# ====================================================================== Reddit
class RedditRSS:
    """비인증 RSS 는 요청 1회 후 ~60초 쿨다운. 응답 헤더의 ratelimit 을 그대로 따른다."""

    def __init__(self):
        self.s = requests.Session()
        self.s.headers['User-Agent'] = UA
        self.wait_until = 0

    def get(self, path):
        url = f'https://www.reddit.com{path}'
        for attempt in range(4):
            gap = self.wait_until - time.time()
            if gap > 0:
                time.sleep(gap)
            r = self.s.get(url, timeout=30)
            rem = float(r.headers.get('x-ratelimit-remaining', '1') or 1)
            reset = float(r.headers.get('x-ratelimit-reset', '0') or 0)
            if r.status_code == 429 or rem < 1:
                self.wait_until = time.time() + max(reset, 5) + 2
            if r.status_code == 200:
                return r.text
            log(f'[reddit] {r.status_code} {path} (재시도 {attempt + 1})')
            if r.status_code != 429:
                self.wait_until = max(self.wait_until, time.time() + 15)
        raise RuntimeError(f'Reddit 요청 실패: {path}')

    @staticmethod
    def entries(xml):
        out = []
        for e in re.findall(r'<entry>(.*?)</entry>', xml, re.S):
            def tag(name):
                m = re.search(rf'<{name}>(.*?)</{name}>', e, re.S)
                return html.unescape(m.group(1)) if m else ''
            link = re.search(r'<link href="([^"]+)"', e)
            author = re.search(r'<name>/u/([^<]+)</name>', e)
            content = re.search(r'<content type="html">(.*?)</content>', e, re.S)
            ts = tag('published') or tag('updated')
            out.append({
                'id': tag('id'), 'title': tag('title'), 'link': link.group(1) if link else '',
                'author': author.group(1) if author else '',
                'html': html.unescape(content.group(1)) if content else '',
                'ts': datetime.fromisoformat(ts).astimezone(KST) if ts else None,
            })
        return out


def _reddit_text(h):
    soup = BeautifulSoup(h, 'html.parser')
    md = soup.select_one('div.md')
    txt = md.get_text(' ', strip=True) if md else ''
    return re.sub(r'\s+', ' ', txt)


def collect_reddit(ws, we):
    rr = RedditRSS()
    posts_xml = rr.get(f'/r/{SUB}/new/.rss?limit=100')
    posts = []
    for e in rr.entries(posts_xml):
        if e['ts'] and ws <= e['ts'] < we and e['author'].lower() not in BOT_AUTHORS:
            body = _reddit_text(e['html'])
            posts.append({
                'id': f"rd:{e['id']}", 'src': 'reddit', 'kind': '게시글', 'ts': iso(e['ts']),
                'post_id': e['id'].replace('t3_', ''), 'title': e['title'], 'text': (e['title'] + ' ' + body).strip(),
                'url': e['link'], 'author': e['author'],
            })
    # 서브레딧 전체 댓글 피드: 100개가 ~30분 분량이라 창 시작 전까지 페이지를 넘긴다
    comments, after, pages = [], None, 0
    while pages < 4:
        q = f'/r/{SUB}/comments/.rss?limit=100' + (f'&after={after}' if after else '')
        ents = rr.entries(rr.get(q))
        pages += 1
        if not ents:
            break
        for e in ents:
            if not e['ts'] or not (ws <= e['ts'] < we) or e['author'].lower() in BOT_AUTHORS:
                continue
            m = re.search(r'/comments/([a-z0-9]+)/', e['link'])
            ptitle = re.sub(r'^/u/\S+ on ', '', e['title'])
            text = _reddit_text(e['html'])
            if not text:
                continue
            comments.append({
                'id': f"rd:{e['id']}", 'src': 'reddit', 'kind': '댓글', 'ts': iso(e['ts']),
                'post_id': m.group(1) if m else '', 'title': ptitle, 'text': text, 'url': e['link'],
                'author': e['author'],
            })
        oldest = min((e['ts'] for e in ents if e['ts']), default=None)
        if oldest is None or oldest < ws:
            break
        after = ents[-1]['id']
    # Hot 순위(동점 시 보조 지표)
    hot_rank = {}
    try:
        for i, e in enumerate(rr.entries(rr.get(f'/r/{SUB}/hot/.rss?limit=50'))):
            hot_rank[e['id'].replace('t3_', '')] = i + 1
    except Exception as e:  # noqa
        log(f'[reddit] hot 실패(무시): {e}')
    return posts, comments, {'hot_rank': hot_rank, 'comment_pages': pages}


# ====================================================================== DC
# DC 는 몰아서 요청하면 IP 단위로 빈 응답(200, 0바이트)을 준다(첫 시험 때 ~100건/25초에서 차단).
# 그래서 ① 목록은 100개/페이지인 PC판, ② 글은 본문+최신 댓글 100개가 한 번에 오는
# 모바일판을 1초 이상 간격으로 한 번씩만 열고, ③ 직전 창 글은 댓글 수가 늘었을 때만 다시 연다.
_DC_PATH = {'board': 'board', 'mgallery': 'mgallery/board', 'mini': 'mini/board'}[DC_TYPE]
DC_LIST = f'https://gall.dcinside.com/{_DC_PATH}/lists/'
DC_VIEW = f'https://gall.dcinside.com/{_DC_PATH}/view/'
DC_M = f'https://m.dcinside.com/board/{DC_ID}'
MUA = ('Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X) AppleWebKit/605.1.15 '
       '(KHTML, like Gecko) Version/17.0 Mobile/15E148 Safari/604.1')
DC_GAP = 1.1          # 요청 간격(초) + 0~0.6초 무작위
DC_MAX_VIEWS = 260    # 한 회차 글 방문 상한


class DCBlocked(Exception):
    pass


class DCClient:
    def __init__(self):
        self.pc = requests.Session()
        self.pc.headers['User-Agent'] = UA
        self.m = requests.Session()
        self.m.headers['User-Agent'] = MUA
        self.last = 0
        self.fails = 0

    def get(self, sess, url, **kw):
        for attempt in range(2):
            gap = DC_GAP + random.random() * 0.6 - (time.time() - self.last)
            if gap > 0:
                time.sleep(gap)
            self.last = time.time()
            try:
                r = sess.get(url, timeout=20, **kw)
                if r.status_code == 200 and len(r.content) > 500:
                    self.fails = 0
                    return r.text
            except requests.RequestException:
                pass
            self.fails += 1
            if self.fails >= 4:
                raise DCBlocked(url)
            time.sleep(20 * (attempt + 1))
        return None


def _dc_txt(node):
    if node is None:
        return ''
    t = node.get_text(' ', strip=True)
    t = re.sub(r'\s*-\s*dc\s*(official\s*)?App\s*$', '', t)
    return re.sub(r'\s+', ' ', t).strip()


def _dc_list_pc(dc, cutoff):
    rows = []
    for page in range(1, 7):
        try:
            h = dc.get(dc.pc, DC_LIST, params={'id': DC_ID, 'list_num': 100, 'page': page})
        except DCBlocked:
            h = None
        if not h:
            return rows or None
        soup = BeautifulSoup(h, 'html.parser')
        oldest = None
        for tr in soup.select('tr.ub-content'):
            no, typ = tr.get('data-no'), tr.get('data-type') or ''
            subj = tr.select_one('.gall_subject')
            subj = subj.get_text(strip=True) if subj else ''
            if not no or not no.isdigit() or 'notice' in typ or subj in ('공지', '설문', 'AD'):
                continue
            try:
                ts = datetime.strptime(tr.select_one('.gall_date').get('title'), '%Y-%m-%d %H:%M:%S').replace(tzinfo=KST)
            except Exception:  # noqa
                continue
            a = tr.select_one('.gall_tit a')
            num = lambda sel: int(re.sub(r'\D', '', tr.select_one(sel).get_text()) or 0)  # noqa
            rn = tr.select_one('.reply_num')
            rows.append({'no': no, 'ts': ts, 'title': a.get_text(' ', strip=True) if a else '', 'subject': subj,
                         'views': num('.gall_count'), 'reco': num('.gall_recommend'),
                         'replies': int(re.sub(r'\D', '', rn.get_text()) or 0) if rn else 0})
            oldest = ts if oldest is None or ts < oldest else oldest
        if oldest is None or oldest < cutoff:
            break
    return rows


def _dc_list_mobile(dc, cutoff, now):
    """모바일 목록(20개/페이지, 시각은 분 단위)."""
    rows = []
    dc.fails = 0
    for page in range(1, 31):
        h = dc.get(dc.m, DC_M, params={'page': page})
        if not h:
            break
        soup = BeautifulSoup(h, 'html.parser')
        oldest = None
        for li in soup.select('ul.gall-detail-lst > li'):
            a = li.select_one('a.lt')
            m = re.search(r'/board/\w+/(\d+)', a['href']) if a else None
            if not m:
                continue
            info = [x.get_text(' ', strip=True) for x in li.select('ul.ginfo > li')]
            tm = next((x for x in info if re.fullmatch(r'\d{1,2}:\d{2}|\d{2}\.\d{2}', x)), None)
            if not tm or ':' not in tm:   # 'MM.DD' = 어제 이전 글
                oldest = cutoff - timedelta(days=1)
                continue
            hh, mm = map(int, tm.split(':'))
            ts = now.replace(hour=hh, minute=mm, second=0, microsecond=0)
            if ts > now + timedelta(minutes=1):
                ts -= timedelta(days=1)
            subj = info[0] if info else ''
            if subj in ('공지', '설문', 'AD'):
                continue
            g = lambda k: int(re.sub(r'\D', '', next((x for x in info if x.startswith(k)), '0')) or 0)  # noqa
            ct = li.select_one('.ct')
            rows.append({'no': m.group(1), 'ts': ts, 'title': _dc_txt(li.select_one('.subjectin')), 'subject': subj,
                         'views': g('조회'), 'reco': g('추천'),
                         'replies': int(re.sub(r'\D', '', ct.get_text()) or 0) if ct else 0})
            oldest = ts if oldest is None or ts < oldest else oldest
        if oldest is None or oldest < cutoff:
            break
    return rows


def _dc_view(dc, no, we):
    """모바일 글 페이지 1회 → (본문, [(댓글no, ts, text)]). 댓글은 최신 100개."""
    h = dc.get(dc.m, f'{DC_M}/{no}')
    if not h:
        return None
    soup = BeautifulSoup(h, 'html.parser')
    body = _dc_txt(soup.select_one('.thum-txtin'))
    cmts = []
    for li in soup.select('ul.all-comment-lst li[no]'):
        d = li.select_one('.date')
        txt = _dc_txt(li.select_one('.txt'))
        if not d or not txt or '삭제된 댓글' in txt:
            continue
        ds = d.get_text(strip=True)
        try:
            if re.match(r'\d{4}\.', ds):
                ts = datetime.strptime(ds[:16], '%Y.%m.%d %H:%M')
            else:
                ts = datetime.strptime(f'{we.year}.{ds[:11]}', '%Y.%m.%d %H:%M')
        except ValueError:
            continue
        cmts.append((li.get('no'), ts.replace(tzinfo=KST), txt))
    return body, cmts


def collect_dc(ws, we):
    dc = DCClient()
    now = datetime.now(KST)
    cutoff = ws - WINDOW  # 직전 창 글에도 이번 창 댓글이 달린다
    src = 'pc'
    rows = _dc_list_pc(dc, cutoff)
    if not rows:
        log('[dc] PC 목록 응답 없음 → 모바일 목록으로 대체')
        src = 'mobile'
        rows = _dc_list_mobile(dc, cutoff, now)
    if not rows:
        raise RuntimeError('DC 목록을 받지 못함(일시 차단 가능성)')
    seen_p = os.path.join(DATA, 'dc_seen.json')
    try:
        seen = json.load(open(seen_p, encoding='utf-8'))
    except Exception:  # noqa
        seen = {}
    uniq = {}
    for x in rows:
        uniq.setdefault(x['no'], x)
    in_win = [x for x in uniq.values() if ws <= x['ts'] < we]
    prev = [x for x in uniq.values() if cutoff <= x['ts'] < ws and x['replies'] > seen.get(x['no'], 0)]
    targets = (sorted(in_win, key=lambda x: -x['replies']) + sorted(prev, key=lambda x: -x['replies']))[:DC_MAX_VIEWS]

    posts, comments, partial, fetched = [], [], False, 0
    in_ids = {x['no'] for x in in_win}
    dc.fails = 0
    for x in targets:
        try:
            v = _dc_view(dc, x['no'], we)
        except DCBlocked:
            partial = True
            log(f'[dc] 연속 빈 응답 — 글 {fetched}/{len(targets)} 방문에서 중단(부분 결과)')
            break
        fetched += 1
        body, cl = v if v else ('', [])
        if v is not None:
            seen[x['no']] = x['replies']
        if x['no'] in in_ids:
            posts.append({
                'id': f"dc:{x['no']}", 'src': 'dc', 'kind': '게시글', 'ts': iso(x['ts']), 'post_id': x['no'],
                'title': x['title'], 'text': (x['title'] + ' ' + body[:1500]).strip(),
                'url': f"{DC_VIEW}?id={DC_ID}&no={x['no']}", 'views': x['views'], 'reco': x['reco'],
                'replies': x['replies'], 'subject': x['subject'],
            })
        for cno, ts, txt in cl:
            if ws <= ts < we:
                comments.append({
                    'id': f"dc:{x['no']}:{cno}", 'src': 'dc', 'kind': '댓글', 'ts': iso(ts), 'post_id': x['no'],
                    'title': x['title'], 'text': txt, 'url': f"{DC_VIEW}?id={DC_ID}&no={x['no']}",
                })
    # 중단 등으로 못 연 창 내 글도 제목으로는 집계한다
    got = {p['post_id'] for p in posts}
    for x in in_win:
        if x['no'] not in got:
            posts.append({'id': f"dc:{x['no']}", 'src': 'dc', 'kind': '게시글', 'ts': iso(x['ts']), 'post_id': x['no'],
                          'title': x['title'], 'text': x['title'], 'url': f"{DC_VIEW}?id={DC_ID}&no={x['no']}",
                          'views': x['views'], 'reco': x['reco'], 'replies': x['replies'], 'subject': x['subject']})
    top_no = max(int(k) for k in uniq)
    keep = {k: v for k, v in seen.items() if k.isdigit() and int(k) > top_no - 20000}
    json.dump(keep, open(seen_p, 'w', encoding='utf-8'))
    return posts, comments, {'list': src, 'views_fetched': fetched, 'views_planned': len(targets),
                             'partial': partial, 'prev_window_revisits': len(prev)}


# ====================================================================== Steam 토론
# 목록은 최근 활동순 15개/페이지, 글 페이지는 오래된 댓글부터 15개/페이지(?ctp=N)이고
# 원글(.forum_op)은 모든 페이지 위에 다시 나온다. 그래서 목록의 댓글 수로 마지막 페이지를
# 계산해 바로 열고, 그 페이지 첫 댓글도 창 안이면 한 페이지씩 앞으로 간다.
# 댓글 속 인용("Originally posted by …")은 남의 말이므로 blockquote 를 지우고 분류한다.
SF_BASE = f'https://steamcommunity.com/app/{APP_ID}/discussions'
SF_FORUMS = CFG.get('steam_forums') or {}  # 비어 있으면 첫 실행 때 게시판 목록을 자동으로 찾는다


def _sf_discover(sf):
    h = sf.get(f'{SF_BASE}/')
    found = {}
    for m in re.finditer(rf'href="https://steamcommunity.com/app/{APP_ID}/discussions/(\d+)/"[^>]*>([^<]+)<', h or ''):
        found.setdefault(m.group(1), html.unescape(m.group(2)).strip())
    return found or {'0': 'General'}
SF_GAP = 1.5
SF_PAGE = 15


class SteamForum:
    def __init__(self):
        self.s = requests.Session()
        self.s.headers['User-Agent'] = UA
        self.s.headers['Accept-Language'] = 'en-US,en;q=0.9'
        self.last = 0

    def get(self, url, **kw):
        for attempt in range(3):
            gap = SF_GAP - (time.time() - self.last)
            if gap > 0:
                time.sleep(gap)
            self.last = time.time()
            try:
                r = self.s.get(url, timeout=20, **kw)
                if r.status_code == 200 and len(r.content) > 1000:
                    return r.text
                log(f'[steamforum] {r.status_code} {url} (재시도 {attempt + 1})')
            except requests.RequestException as e:
                log(f'[steamforum] {e} (재시도 {attempt + 1})')
            time.sleep(10 * (attempt + 1))
        return None


def _sf_ts(node):
    t = node.select_one('.commentthread_comment_timestamp[data-timestamp]') if node is not None else None
    return datetime.fromtimestamp(int(t['data-timestamp']), KST) if t else None


def _sf_text(node):
    if node is None:
        return ''
    for b in node.select('blockquote'):
        b.decompose()
    return re.sub(r'\s+', ' ', node.get_text(' ', strip=True)).strip()


def collect_steamforum(ws, we):
    sf = SteamForum()
    topics = []
    forums = SF_FORUMS or _sf_discover(sf)
    for fid, fname in forums.items():
        for page in range(1, 6):
            h = sf.get(f'{SF_BASE}/{fid}/', params={'fp': page} if page > 1 else None)
            if not h:
                if page == 1 and fid == next(iter(forums)):
                    raise RuntimeError('Steam 일반 토론 목록을 받지 못함')
                break
            soup = BeautifulSoup(h, 'html.parser')
            rows = soup.select('.forum_topic')
            older = False
            for t in rows:
                lp = t.select_one('.forum_topic_lastpost')
                ts = datetime.fromtimestamp(int(lp['data-timestamp']), KST) if lp and lp.get('data-timestamp') else None
                sticky = 'sticky' in (t.get('class') or [])
                if ts is None or ts < ws:
                    if not sticky:  # 고정글은 맨 위라 활동순 정렬과 무관
                        older = True
                    continue
                a = t.select_one('a.forum_topic_overlay')
                rc = re.sub(r'\D', '', t.select_one('.forum_topic_reply_count').get_text()) if t.select_one('.forum_topic_reply_count') else ''
                topics.append({'gid': t.get('data-gidforumtopic'), 'url': a['href'] if a else '', 'forum': fname,
                               'replies': int(rc or 0), 'sticky': sticky,
                               'title': _sf_text(t.select_one('.forum_topic_name')).replace('PINNED: ', '')})
            if older or not rows:
                break

    posts, comments = [], []
    for tp in topics:
        if not tp['url']:
            continue
        page = max(1, -(-tp['replies'] // SF_PAGE))
        op_done = False
        for _ in range(4):
            h = sf.get(tp['url'], params={'ctp': page} if page > 1 else None)
            if not h:
                break
            soup = BeautifulSoup(h, 'html.parser')
            op = soup.select_one('.forum_op')
            if op is not None and not op_done:
                op_done = True
                ots = _sf_ts(op)
                title = _sf_text(op.select_one('.topic')) or tp['title']
                tp['title'] = title
                if ots and ws <= ots < we:
                    body = _sf_text(op.select_one('.content'))
                    posts.append({'id': f"sf:{tp['gid']}", 'src': 'steamforum', 'kind': '게시글', 'ts': iso(ots),
                                  'post_id': tp['gid'], 'title': title, 'text': (title + ' ' + body[:1500]).strip(),
                                  'url': tp['url'], 'replies': tp['replies'], 'forum': tp['forum']})
            cl = soup.select('.commentthread_comment')
            first_in = False
            for i, c in enumerate(cl):
                ts = _sf_ts(c)
                if ts is None:
                    continue
                if i == 0 and ts >= ws:
                    first_in = True
                if not (ws <= ts < we):
                    continue
                text = _sf_text(c.select_one('.commentthread_comment_text'))
                if not text:
                    continue
                cid = (c.get('id') or '').replace('comment_', '')
                comments.append({'id': f"sf:{tp['gid']}:{cid}", 'src': 'steamforum', 'kind': '댓글', 'ts': iso(ts),
                                 'post_id': tp['gid'], 'title': tp['title'], 'text': text,
                                 'url': f"{tp['url']}#c{cid}", 'post_url': tp['url'], 'replies': tp['replies']})
            if first_in and page > 1:
                page -= 1
                continue
            break
    return posts, comments, {'topics_active': len(topics),
                             'reply_total': {t['gid']: t['replies'] for t in topics}}


# ====================================================================== 라벨 / 학습
def _fernet(key=None):
    from cryptography.fernet import Fernet
    key = key or os.environ.get('LABELS_KEY') or open(LABELS_KEY_FILE, encoding='utf-8').read().strip()
    return Fernet(key.encode())


def export_labels(out_path=LABELS_ENC):
    """로컬 → 클라우드: 엑셀 라벨(판단 이유 포함)을 암호화해 labels.enc 로. 키가 없으면 새로 만든다(OneDrive 밖)."""
    if not os.path.exists(LABELS_KEY_FILE):
        from cryptography.fernet import Fernet
        os.makedirs(os.path.dirname(LABELS_KEY_FILE), exist_ok=True)
        open(LABELS_KEY_FILE, 'w', encoding='utf-8').write(Fernet.generate_key().decode())
    labeled, files = load_labels()
    blob = json.dumps({'rows': labeled, 'files': files}, ensure_ascii=False).encode('utf-8')
    open(out_path, 'wb').write(_fernet().encrypt(blob))
    return len(labeled)


def load_labels():
    """라벨링 폴더의 모든 검토 엑셀에서 '내 판단' 이 채워진 행을 읽는다(mtime 캐시).
    클라우드에서는 엑셀 대신 labels.enc(암호화, 키는 GitHub 비밀값 LABELS_KEY)를 읽는다."""
    if CLOUD:
        try:
            d = json.loads(_fernet().decrypt(open(LABELS_ENC, 'rb').read()).decode('utf-8'))
            return d['rows'], d['files']
        except Exception as e:  # noqa
            log(f'[label] labels.enc 읽기 실패 → 라벨 없이 진행: {e}')
            return [], []
    from openpyxl import load_workbook
    cache_p = os.path.join(DATA, 'labels_cache.json')
    try:
        cache = json.load(open(cache_p, encoding='utf-8'))
    except Exception:  # noqa
        cache = {}
    out_cache = {}
    for fn in sorted(os.listdir(LABEL_DIR)):
        if not fn.endswith('.xlsx') or fn.startswith('~$'):
            continue
        p = os.path.join(LABEL_DIR, fn)
        mt = os.path.getmtime(p)
        if fn in cache and cache[fn]['mtime'] == mt:
            out_cache[fn] = cache[fn]
            continue
        rows = []
        try:
            wb = load_workbook(p, read_only=True, data_only=True)
            ws = wb['검토']
            header = None
            for r in ws.iter_rows(values_only=True):
                if header is None:
                    if r and '내 판단' in [str(c).strip() if c else '' for c in r]:
                        header = [str(c).strip() if c else '' for c in r]
                    continue
                d = dict(zip(header, r))
                lab = str(d.get('내 판단') or '').strip()
                lab = {'긍': '긍정', '부': '부정', '중': '중립', 'p': '긍정', 'n': '부정', '0': '중립'}.get(lab, lab)
                if lab in S.LABELS and d.get('ID'):
                    rows.append({'id': d['ID'], 'text': str(d.get('본문') or ''), 'label': lab,
                                 'pred': str(d.get('프로그램 판단') or ''), 'src': d.get('출처'),
                                 'reason_type': str(d.get('이유 유형') or '').strip(),
                                 'reason': str(d.get('이유 설명') or '').strip(), 'file': fn})
            wb.close()
        except Exception as e:  # noqa
            log(f'[label] {fn} 읽기 실패(엑셀이 열려 있으면 다음 회차에 반영): {e}')
            if fn in cache:
                out_cache[fn] = cache[fn]
            continue
        out_cache[fn] = {'mtime': mt, 'rows': rows}
    json.dump(out_cache, open(cache_p, 'w', encoding='utf-8'), ensure_ascii=False)
    labeled = {}
    files = []
    for fn, v in out_cache.items():
        files.append({'file': fn, 'labeled': len(v['rows']),
                      'agree': sum(1 for r in v['rows'] if r['pred'] == r['label'])})
        for r in v['rows']:
            labeled[r['id']] = r
    return list(labeled.values()), files


def build_model():
    labeled, files = load_labels()
    S.set_user_lexicon(user_lexicon(labeled))
    # 학습기: 라벨이 바뀔 때만 교차검증으로 모델을 다시 고르고(수 초), 결과는 지문별로 캐시한다.
    # 로지스틱 회귀가 '사전만'보다 CV 정확도가 높을 때만 채택 — 학습이 성능을 깎지 못하게(2026-10-06 NB 사고).
    sig = label_signature(labeled) + ':' + S.LEXICON_VERSION
    cache_p = os.path.join(DATA, 'model_select.json')
    try:
        cache = json.load(open(cache_p, encoding='utf-8'))
    except Exception:  # noqa
        cache = {}
    cached = cache.get('sel') if cache.get('sig') == sig else None
    nb = S.Learner(labeled, cached) if labeled else None
    if nb is not None and nb.cv and not cached:
        try:
            json.dump({'sig': sig, 'sel': nb.cv}, open(cache_p, 'w', encoding='utf-8'), ensure_ascii=False)
        except OSError:
            pass
    return nb, labeled, files


_QUOTE = re.compile(r"[\"'“”‘’「」『』]([^\"'“”‘’「」『』\n]{1,20})[\"'“”‘’「」『』]")
_REASON_W = {'긍정': 2.0, '부정': -2.0, '중립': 0.0}


def user_lexicon(labeled):
    """'이유 설명'에 따옴표로 적은 표현 중 실제 본문에 있는 것을 사전에 넣는다.
    긍정 +2 / 부정 -2 / 중립 0(그 표현은 판단 근거 아님). 여러 라벨이 엇갈리면 다수결, 동률이면 0.
    모든 이유는 data/label_reasons.jsonl 로 모아 둔다 — 사전·규칙을 손볼 때 먼저 읽는 자료."""
    votes, out_rows = {}, []
    for r in labeled:
        terms, unmatched = [], []
        for m in _QUOTE.finditer(r.get('reason') or ''):
            t = m.group(1).strip().lower()
            if len(t) < 1:
                continue
            (terms if t in (r['text'] or '').lower() else unmatched).append(t)
        for t in terms:
            votes.setdefault(t, []).append(r['label'])
        if r.get('reason') or r.get('reason_type'):
            out_rows.append({'id': r['id'], 'label': r['label'], 'pred': r.get('pred'), 'src': r.get('src'),
                             'reason_type': r.get('reason_type'), 'reason': r.get('reason'), 'terms': terms,
                             'unmatched_terms': unmatched, 'text': (r['text'] or '')[:500], 'file': r.get('file')})
    lex = {}
    for t, labs in votes.items():
        c = {l: labs.count(l) for l in set(labs)}
        top = max(c.values())
        winners = [l for l, n in c.items() if n == top]
        lex[t] = _REASON_W[winners[0]] if len(winners) == 1 else 0.0
    try:
        with open(os.path.join(DATA, 'label_reasons.jsonl'), 'w', encoding='utf-8') as f:
            for row in out_rows:
                f.write(json.dumps(row, ensure_ascii=False) + '\n')
        json.dump(lex, open(os.path.join(DATA, 'user_lexicon.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    except OSError:
        pass
    return lex


def label_signature(labeled):
    """라벨 묶음의 지문. 이게 바뀌면(추가·수정·삭제) 과거 수집분 전체를 다시 분류한다."""
    import hashlib
    body = '\n'.join(f"{r['id']}\t{r['label']}\t{r.get('reason_type', '')}\t{r.get('reason', '')}"
                     for r in sorted(labeled, key=lambda r: str(r['id'])))
    return hashlib.md5((body + S.LEXICON_VERSION).encode('utf-8')).hexdigest()


def classify_item(it, nb, user):
    """사용자가 엑셀에서 직접 판단한 항목은 그 판단이 최우선."""
    lab = user.get(it['id'])
    if lab:
        it.update({'label': lab, 'score': {'긍정': 3, '중립': 0, '부정': -3}[lab], 'hits': ['사용자 판단'],
                   'ambiguous': False, 'method': '사용자'})
        return it
    it.update(S.classify(it['text'], nb))
    st = ai_state()
    if st.get('enabled') and it.get('ai_label'):
        import llm
        lex = it['label']
        it['lex_label'] = lex
        it['label'] = llm.combine(st['variant'], lex, it['ai_label'], it.get('ai_conf'))
        it['method'] = f"AI조합({st['variant']})"
        conf = it.get('ai_conf')
        it['ambiguous'] = conf == 'l' or (bool(it['ambiguous']) and conf != 'h')
    return it


# ---------------------------------------------------------------- AI 분류 (자동 재시험·조건부 사용)
AI_STATE_P = os.path.join(DATA, 'ai_mode.json')
AI_RETEST_EVERY = int(CFG.get('ai_retest_every_labels', 100))
_ai_state = None


def ai_state():
    global _ai_state
    if _ai_state is None:
        try:
            _ai_state = json.load(open(AI_STATE_P, encoding='utf-8'))
        except Exception:  # noqa
            _ai_state = {'enabled': False, 'variant': None, 'last_eval_n': 0, 'history': []}
    return _ai_state


def save_ai_state(st):
    global _ai_state
    _ai_state = st
    tmp = AI_STATE_P + '.tmp'
    json.dump(st, open(tmp, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    os.replace(tmp, AI_STATE_P)


def labeled_with_context(labeled):
    """라벨 행에 게시글 제목·유형을 붙인다(AI 예시·시험 입력용)."""
    want = {r['id'] for r in labeled}
    ctx = {}
    for fn in item_files():
        for it in read_jsonl(fn):
            if it['id'] in want:
                ctx[it['id']] = it
    return [{**r, 'title': ctx.get(r['id'], {}).get('title', ''), 'kind': ctx.get(r['id'], {}).get('kind', '댓글')}
            for r in labeled]


def apply_ai(items, labeled, user):
    """AI 모드가 켜져 있을 때 이번 회차 항목에 AI 라벨을 붙인다(사용자 판단 항목 제외). 실패분은 사전 그대로."""
    st = ai_state()
    if not st.get('enabled'):
        return
    import llm
    llm.API_DAILY_BUDGET = float(CFG.get('llm_api_daily_usd', llm.API_DAILY_BUDGET))
    todo = [it for it in items if it['id'] not in user]
    t0 = time.time()
    res, errs = llm.classify_items(todo, GAME, labeled_with_context(labeled), tag='run')
    got = 0
    for it, r in zip(todo, res):
        if r:
            it['ai_label'], it['ai_conf'] = r
            got += 1
    log(f"[ai] {got}/{len(todo)}건 AI 분류 ({time.time() - t0:.0f}초, 조합 {st['variant']})"
        + (f" · 오류 {errs[0][:120]}" if errs else ''))


def maybe_retest_ai(labeled):
    """라벨이 지난 시험보다 AI_RETEST_EVERY 건 이상 늘었으면 5-fold 로 다시 시험한다(구독 경로만).
    가장 좋은 조합이 '사전 + 5%p' 이상이면 그 조합으로 AI 분류를 켠다(사용자 합의 기준, 2026-10-07)."""
    st = ai_state()
    if len(labeled) < st.get('last_eval_n', 0) + AI_RETEST_EVERY:
        return
    import llm
    log(f"[ai-eval] 라벨 {len(labeled)}건 (지난 시험 {st.get('last_eval_n', 0)}건) → AI 재시험 시작")
    t0 = time.time()
    r = llm.evaluate(labeled_with_context(labeled), GAME, lambda t: S.classify(t, None)['label'])
    r['ts'] = iso(datetime.now(KST))
    r['seconds'] = round(time.time() - t0)
    st.setdefault('history', []).append(r)
    if not r.get('ok'):
        log(f"[ai-eval] 무효(실패 {r.get('failed')}건, 구독 한도 등) — 다음 회차에 다시 시도: {r.get('errors')}")
        save_ai_state(st)
        return
    st['last_eval_n'] = r['n']
    if r['passed']:
        st['enabled'], st['variant'], st['enabled_at'] = True, r['best'], r['ts']
        log(f"[ai-eval] 통과 — {llm.VARIANTS[r['best']]} {r['best_score']}% ≥ 기준 {r['threshold']}% → AI 분류 켬")
    else:
        if st.get('enabled'):  # 켜진 뒤 다시 미달이면 끈다
            st['enabled'], st['disabled_at'] = False, r['ts']
        log(f"[ai-eval] 미달 — 최고 {llm.VARIANTS[r['best']]} {r['best_score']}% < 기준 {r['threshold']}% (사전 {r['lexicon']}%) → 사전 유지")
    save_ai_state(st)


class DataLock:
    """수집(community)과 재계산(rescore)이 items·runs 파일을 동시에 고쳐 쓰지 않도록."""
    path = os.path.join(DATA, '.lock')

    def __enter__(self):
        t0 = time.time()
        while True:
            try:
                fd = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
                os.write(fd, str(os.getpid()).encode())
                os.close(fd)
                return self
            except FileExistsError:
                try:
                    if time.time() - os.path.getmtime(self.path) > 40 * 60:  # 죽은 프로세스가 남긴 잠금
                        os.remove(self.path)
                        continue
                except FileNotFoundError:
                    continue
                if time.time() - t0 > 20 * 60:
                    raise RuntimeError('데이터 잠금 대기 시간 초과')
                time.sleep(5)

    def __exit__(self, *a):
        try:
            os.remove(self.path)
        except FileNotFoundError:
            pass


def model_state():
    try:
        return json.load(open(os.path.join(DATA, 'model_state.json'), encoding='utf-8'))
    except Exception:  # noqa
        return {}


# ====================================================================== 30분 창 분석
def current_window(now=None):
    now = now or datetime.now(KST)
    we = now.replace(second=0, microsecond=0, minute=(now.minute // 30) * 30)
    return we - WINDOW, we


def neg_topic_counts(items):
    """부정으로 분류된 항목(게시글+댓글)의 주제 분포 {주제: 건수} — 대시보드 '부정 반응 성격 추이'용."""
    from collections import Counter
    return dict(Counter(S.topic_of(it.get('text', '')) for it in items if it.get('label') == '부정'))


def ensure_neg_topics():
    """neg_topics 가 없는 과거 회차를 그 창의 항목 파일로 채운다(기능 추가 전 회차, 1회성)."""
    runs_p = os.path.join(DATA, 'runs.jsonl')
    runs = read_jsonl(runs_p)
    todo = [r for r in runs if any(o.get('ok') and 'neg_topics' not in o for o in r['sources'].values())]
    if not todo:
        return
    with DataLock():
        runs = read_jsonl(runs_p)
        n = 0
        for r in runs:
            if not any(o.get('ok') and 'neg_topics' not in o for o in r['sources'].values()):
                continue
            its = read_jsonl(item_path(r['slot_end']))
            for src, o in r['sources'].items():
                if o.get('ok') and 'neg_topics' not in o:
                    o['neg_topics'] = neg_topic_counts([x for x in its if x['src'] == src])
            n += 1
        tmp = runs_p + '.tmp'
        with open(tmp, 'w', encoding='utf-8') as f:
            for r in runs:
                f.write(json.dumps(r, ensure_ascii=False) + '\n')
        os.replace(tmp, runs_p)
    log(f'[neg-topics] 과거 {n}개 창의 부정 반응 주제 분포를 채움')


def summarize(items):
    c = {'긍정': 0, '중립': 0, '부정': 0}
    for it in items:
        c[it['label']] += 1
    n = sum(c.values())
    polar = c['긍정'] + c['부정']
    return {
        'n': n, 'pos': c['긍정'], 'neu': c['중립'], 'neg': c['부정'],
        'pos_pct': round(c['긍정'] / n * 100, 1) if n else None,
        'neg_pct': round(c['부정'] / n * 100, 1) if n else None,
        'pos_share': round(c['긍정'] / polar * 100, 1) if polar else None,  # 긍정/(긍정+부정)
        'posts': sum(1 for it in items if it['kind'] == '게시글'),
        'comments': sum(1 for it in items if it['kind'] == '댓글'),
    }


def top3(src, posts, comments, meta):
    by_post = {}
    for c in comments:
        by_post.setdefault(c['post_id'], []).append(c)
    cands = {}
    for p in posts:
        cands[p['post_id']] = {'post_id': p['post_id'], 'title': p['title'], 'url': p['url'], 'item': p}
    if src in ('reddit', 'steamforum'):
        # Reddit·Steam 토론은 조회수를 공개하지 않는다 → 30분간 새 댓글 수(반응량)로 대체.
        # 동점이면 Reddit 은 Hot 순위, Steam 은 누적 댓글 수.
        for pid, cl in by_post.items():
            url = (f'https://www.reddit.com/r/{SUB}/comments/{pid}/' if src == 'reddit' else cl[0].get('post_url', cl[0]['url']))
            cands.setdefault(pid, {'post_id': pid, 'title': cl[0]['title'], 'url': url, 'item': None})
        hot = meta.get('hot_rank', {})
        tot = meta.get('reply_total', {})
        ranked = sorted(cands.values(), key=lambda c: (-len(by_post.get(c['post_id'], [])),
                                                       hot.get(c['post_id'], 999), -tot.get(c['post_id'], 0)))
    else:
        ranked = sorted(cands.values(), key=lambda c: -(c['item'] or {}).get('views', 0))
    out = []
    for c in ranked[:3]:
        cl = by_post.get(c['post_id'], [])
        it = c['item']
        text = (it['text'] if it else c['title'])
        out.append({
            'post_id': c['post_id'], 'title': c['title'][:140], 'url': c['url'],
            'label': it['label'] if it else S.classify(c['title'])['label'],
            'topic': S.topic_of(text + ' ' + ' '.join(x['text'] for x in cl[:30])),
            'views': it.get('views') if it else None, 'reco': it.get('reco') if it else None,
            'win_comments': len(cl), 'hot_rank': meta.get('hot_rank', {}).get(c['post_id']),
            'total_replies': meta.get('reply_total', {}).get(c['post_id']),
        })
        _t3_comment_counts(out[-1], cl)
    return out


def _t3_comment_counts(t, cl):
    cs = summarize(cl) if cl else None
    t['cmt_pos'] = cs['pos'] if cs else 0
    t['cmt_neg'] = cs['neg'] if cs else 0
    t['cmt_neu'] = cs['neu'] if cs else 0


def _pid_from_url(src, url):
    """post_id 를 저장하기 전에 만든 Top3 항목용(2026-10-06 초기 회차)."""
    pat = {'reddit': r'/comments/([a-z0-9]+)/', 'dc': r'[?&]no=(\d+)', 'steamforum': r'/discussions/\d+/(\d+)/'}[src]
    m = re.search(pat, url or '')
    return m.group(1) if m else None


def write_review_xlsx(items, ws, we, run_stats):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    rnd = random.Random(iso(we))
    pool = [it for it in items if it['ambiguous']
            and len(re.sub(r'https?://\S+', '', it['text']).strip()) >= 4]  # 링크만 있는 글 제외
    # 출처별 균형: 30개를 출처 수로 나눠 뽑고, 모자란 출처 몫은 다른 출처로 채운다
    by = {src: [x for x in pool if x['src'] == src] for src in SOURCES}
    for v in by.values():
        rnd.shuffle(v)
    share = REVIEW_N // len(SOURCES)
    pick = [x for v in by.values() for x in v[:share]]
    rest = [x for v in by.values() for x in v[share:]]
    rnd.shuffle(rest)
    pick += rest[:REVIEW_N - len(pick)]
    pick.sort(key=lambda x: (x['src'], x['ts']))

    wb = Workbook()
    sh = wb.active
    sh.title = '검토'
    sh['A1'] = f"{GAME} 반응 검토 — {we.strftime('%Y-%m-%d')} {ws.strftime('%H:%M')}~{we.strftime('%H:%M')} 창, 판단이 애매한 {len(pick)}건(무작위)"
    sh['A1'].font = Font(bold=True, size=12)
    sh['A2'] = ("작성법: ① '내 판단'에서 긍정 / 부정 / 중립을 고릅니다(필수). "
                "② 왜 그렇게 봤는지 '이유 유형'(목록)과 '이유 설명'(자유)에 적어 주세요 — 특히 부정일 때. "
                "설명에 판단의 근거가 된 표현을 따옴표로 적으면(예: '창렬'이라는 말, \"숙제 같다\") 그 표현이 사전에 추가되어 "
                "비슷한 글을 같은 방향으로 판단합니다. ③ 저장 후 이 폴더에 그대로 두면 다음 실행 때 과거 데이터까지 다시 계산됩니다. "
                "중립 = 게임에 대한 호불호가 드러나지 않음(질문·정보·잡담).")
    sh['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    sh.merge_cells('A2:N2')
    sh.row_dimensions[2].height = 62
    # (열 이름, 너비) — 학습 로더는 열 '이름'으로 읽으므로 순서를 바꿔도 된다
    spec = [('No', 5), ('내 판단', 11), ('이유 유형', 16), ('이유 설명', 34), ('출처', 10), ('유형', 7),
            ('본문', 60), ('번역(한국어)', 50), ('게시글 제목', 30), ('프로그램 판단', 12), ('점수', 7),
            ('근거 키워드', 26), ('링크', 8), ('ID', 18), ('작성시각', 17)]
    col = {name: i for i, (name, _) in enumerate(spec, 1)}
    hdr_row = 4
    fill = PatternFill('solid', fgColor='1F3A5F')
    thin = Side(style='thin', color='C8CDD3')
    for name, w in spec:
        cell = sh.cell(row=hdr_row, column=col[name], value=name)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        sh.column_dimensions[cell.column_letter].width = w
    sh.cell(row=hdr_row, column=col['내 판단']).fill = PatternFill('solid', fgColor='C0392B')
    for name in ('이유 유형', '이유 설명'):
        sh.cell(row=hdr_row, column=col[name]).fill = PatternFill('solid', fgColor='D35400')
    dv = DataValidation(type='list', formula1='"긍정,부정,중립"', allow_blank=True, showDropDown=False)
    dv.promptTitle = '내 판단'
    dv.prompt = '긍정 / 부정 / 중립'
    sh.add_data_validation(dv)
    dv_r = DataValidation(type='list', formula1='"' + ','.join(REASON_TYPES) + '"', allow_blank=True, showDropDown=False)
    dv_r.promptTitle = '이유 유형(선택)'
    dv_r.prompt = '목록에 없으면 비우고 이유 설명에 적어 주세요'
    sh.add_data_validation(dv_r)
    inp = PatternFill('solid', fgColor='FFF4CC')
    inp2 = PatternFill('solid', fgColor='FDEBD0')
    wrap_cols = {col[n] for n in ('이유 설명', '본문', '번역(한국어)', '게시글 제목', '근거 키워드')}
    for k, it in enumerate(pick, 1):
        r = hdr_row + k
        ko = TR.to_ko(it['text'])          # 한국어 글이면 '' (번역 불필요)
        title = it.get('title', '')[:200]
        title_ko = TR.to_ko(title, 300) if title and title != it['text'][:len(title)] else ''
        if title_ko:
            title = f'{title}\n({title_ko})'
        vals = {'No': k, '출처': SRC_NAME[it['src']], '유형': it['kind'], '본문': it['text'][:2000],
                '번역(한국어)': ko or None, '게시글 제목': title, '프로그램 판단': it['label'], '점수': it['score'],
                '근거 키워드': ', '.join(it['hits'])[:200], '링크': '열기', 'ID': it['id'],
                '작성시각': it['ts'].replace('T', ' ')}
        for name, _ in spec:
            cell = sh.cell(row=r, column=col[name], value=vals.get(name))
            cell.alignment = Alignment(wrap_text=col[name] in wrap_cols, vertical='top')
            cell.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        sh.cell(row=r, column=col['내 판단']).fill = inp
        sh.cell(row=r, column=col['내 판단']).alignment = Alignment(horizontal='center', vertical='top')
        sh.cell(row=r, column=col['이유 유형']).fill = inp2
        sh.cell(row=r, column=col['이유 설명']).fill = inp2
        link = sh.cell(row=r, column=col['링크'])
        link.hyperlink = it['url']
        link.font = Font(color='1F6FEB', underline='single')
        dv.add(sh.cell(row=r, column=col['내 판단']))
        dv_r.add(sh.cell(row=r, column=col['이유 유형']))
        sh.row_dimensions[r].height = min(160, 15 * (1 + max(len(it['text']) // 75, len(ko) // 40)))
    sh.freeze_panes = sh.cell(row=hdr_row + 1, column=col['출처'])
    sh.auto_filter.ref = f"A{hdr_row}:{sh.cell(row=hdr_row, column=len(spec)).column_letter}{hdr_row + len(pick)}"
    TR.save()

    info = wb.create_sheet('이번 창 요약')
    info.append(['출처', '전체', '게시글', '댓글', '긍정', '중립', '부정', '긍정%', '부정%', '긍정/(긍+부)%'])
    for src, st in run_stats.items():
        if st:
            info.append([SRC_NAME[src], st['n'], st['posts'], st['comments'], st['pos'],
                         st['neu'], st['neg'], st['pos_pct'], st['neg_pct'], st['pos_share']])
    fn = f"검토_{we.strftime('%Y-%m-%d_%H%M')}.xlsx"
    p = os.path.join(LABEL_DIR, fn)
    wb.save(p)
    return fn, len(pick)


def run_community(now=None):
    with DataLock():
        return _run_community(now)


def _run_community(now=None):
    ws, we = current_window(now)
    runs_p = os.path.join(DATA, 'runs.jsonl')
    if any(r['slot_end'] == iso(we) for r in read_jsonl(runs_p)):
        log(f'[community] {iso(we)} 창은 이미 처리됨 — 건너뜀')
        build_dashboard()
        return 0
    log(f'[community] 창 {ws:%m-%d %H:%M} ~ {we:%H:%M} 수집 시작')
    nb, labeled, _ = build_model()
    user = {r['id']: r['label'] for r in labeled}
    if label_signature(labeled) != model_state().get('sig'):
        log(f'[community] 라벨 변경 감지({len(labeled)}건) → 과거 수집분 전체 재계산')
        _rescore(nb, labeled)

    results, errors = {}, {}

    def worker(name, fn):
        try:
            results[name] = fn(ws, we)
        except Exception as e:  # noqa
            errors[name] = f'{type(e).__name__}: {e}'
            log(f'[{name}] 실패: {e}\n{traceback.format_exc()}')

    collectors = {'reddit': collect_reddit, 'dc': collect_dc, 'steamforum': collect_steamforum}
    th = [threading.Thread(target=worker, args=(src, collectors[src])) for src in SOURCES]
    for t in th:
        t.start()
    for t in th:
        t.join()

    if not results:
        # 전 소스 실패 = 네트워크 문제(와이파이 전환·로그인 페이지 등, 2026-10-08 08:31 실측).
        # 창을 '처리됨'으로 남기지 않고 실패로 끝내야 스케줄러 재시도(5분 간격 2회)가 이 창을 다시 수집한다.
        log(f'[community] 모든 소스 실패 → 이 창은 기록하지 않고 재시도에 맡김')
        return 1

    try:
        apply_ai([it for r in results.values() for it in r[0] + r[1]], labeled, user)
    except Exception as e:  # noqa
        log(f'[ai] 실패(사전으로 진행): {e}')

    all_items, src_out, stats = [], {}, {}
    for src in SOURCES:
        if src not in results:
            src_out[src] = {'ok': False, 'error': errors.get(src)}
            stats[src] = None
            continue
        posts, comments, meta = results[src]
        for it in posts + comments:
            classify_item(it, nb, user)
            it['slot'] = iso(we)
        st = summarize(posts + comments)
        stats[src] = st
        src_out[src] = {'ok': True, **st, 'top3': top3(src, posts, comments, meta),
                        'neg_topics': neg_topic_counts(posts + comments),
                        'meta': {k: v for k, v in meta.items() if k not in ('hot_rank', 'reply_total')}}
        all_items += posts + comments
        log(f"[{src}] 게시글 {st['posts']} · 댓글 {st['comments']} → 긍정 {st['pos_pct']}% / 부정 {st['neg_pct']}% {src_out[src]['meta']}")

    items_p = item_path(we)
    os.makedirs(os.path.dirname(items_p), exist_ok=True)
    with open(items_p, 'w', encoding='utf-8') as f:
        for it in all_items:
            f.write(json.dumps(it, ensure_ascii=False) + '\n')

    # 검토 엑셀은 REVIEW_EVERY 분마다(기본 60분). 수집·분석은 30분마다 그대로이고,
    # 표본은 그 시간 동안 모인 모든 창(예: 16:00~17:00 = 30분 창 2개)의 애매한 항목에서 뽑는다.
    review_file, review_n = None, 0
    minutes = we.hour * 60 + we.minute
    if all_items and REVIEW_EVERY > 0 and minutes % REVIEW_EVERY == 0:  # 0 이면 검토 엑셀 생성 중단
        try:
            rs = we - timedelta(minutes=REVIEW_EVERY)
            slots = {iso(rs + WINDOW * k) for k in range(1, REVIEW_EVERY // 30)}  # 이번 창 제외 이전 창들
            prev_items = [it for sl in sorted(slots) for it in read_jsonl(item_path(sl))]
            pool = prev_items + all_items
            hour_stats = {src: summarize([x for x in pool if x['src'] == src]) if stats.get(src) else None
                          for src in SOURCES}
            review_file, review_n = write_review_xlsx(pool, rs, we, hour_stats)
            log(f'[review] {review_file} ({review_n}건, {rs:%H:%M}~{we:%H:%M} 표본)')
        except Exception as e:  # noqa
            log(f'[review] 엑셀 실패: {e}')

    append_jsonl(runs_p, {
        'slot_start': iso(ws), 'slot_end': iso(we), 'collected_at': iso(datetime.now(KST)),
        'model': nb.n if nb and nb.usable else 0, 'review_file': review_file, 'review_n': review_n,
        'sources': src_out,
    })
    try:
        maybe_retest_ai(labeled)
    except Exception as e:  # noqa
        log(f'[ai-eval] 오류: {e}')
    build_dashboard()
    return 1 if errors else 0


# ====================================================================== 재분류
def run_rescore():
    with DataLock():
        nb, labeled, _ = build_model()
        _rescore(nb, labeled)
    build_dashboard()


def _rescore(nb, labeled):
    """모든 수집 항목을 현재 모델(사전+학습+사용자 판단)로 다시 분류하고
    창별 비중과 Top3 논조·댓글 반응까지 다시 계산한다. 수집 당시 수치는 orig 로 한 번만 보존."""
    t0 = time.time()
    user = {r['id']: r['label'] for r in labeled}
    by_slot, cmts, post_label = {}, {}, {}
    n_items = 0
    for p in item_files():
        its = read_jsonl(p)
        for it in its:
            classify_item(it, nb, user)
            by_slot.setdefault((it['slot'], it['src']), []).append(it)
            if it['kind'] == '게시글':
                post_label[(it['src'], it['post_id'])] = it['label']
            else:
                cmts.setdefault((it['slot'], it['src'], it['post_id']), []).append(it)
        n_items += len(its)
        tmp = p + '.tmp'
        with open(tmp, 'w', encoding='utf-8') as f:
            for it in its:
                f.write(json.dumps(it, ensure_ascii=False) + '\n')
        os.replace(tmp, p)
    runs_p = os.path.join(DATA, 'runs.jsonl')
    runs = read_jsonl(runs_p)
    for r in runs:
        for src, so in r['sources'].items():
            its = by_slot.get((r['slot_end'], src))
            if not (so.get('ok') and its):
                continue
            if 'orig' not in so:
                so['orig'] = {k: so.get(k) for k in ('pos', 'neu', 'neg', 'pos_pct', 'neg_pct', 'pos_share')}
            so.update(summarize(its))
            so['neg_topics'] = neg_topic_counts(its)
            for t in so.get('top3') or []:
                pid = t.get('post_id') or _pid_from_url(src, t.get('url'))
                t['post_id'] = pid
                if 'label_orig' not in t:
                    t['label_orig'] = t['label']
                t['label'] = post_label.get((src, pid)) or S.classify(t['title'], nb)['label']
                _t3_comment_counts(t, cmts.get((r['slot_end'], src, pid), []))
        r['model'] = nb.n if nb and nb.usable else 0
    tmp = runs_p + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        for r in runs:
            f.write(json.dumps(r, ensure_ascii=False) + '\n')
    os.replace(tmp, runs_p)
    state = {'sig': label_signature(labeled), 'labeled': len(labeled), 'user_applied': len(user),
             'nb_used': bool(nb and nb.usable), 'model_kind': nb.kind if nb else 'lexicon', 'rescored_at': iso(datetime.now(KST)),
             'items': n_items, 'runs': len(runs), 'seconds': round(time.time() - t0, 1)}
    json.dump(state, open(os.path.join(DATA, 'model_state.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    log(f"[rescore] 라벨 {len(labeled)}건 · 모델 {state['model_kind']} (CV {(nb.cv or {}).get('model') if nb else '-'}%) → "
        f"{n_items:,}개 항목 · {len(runs)}개 창 재계산 ({state['seconds']}초)")


def build_dashboard():
    import dashboard
    try:
        ensure_neg_topics()
    except Exception as e:  # noqa
        log(f'[neg-topics] 채우기 실패: {e}')
    nb, labeled, files = build_model()
    runs = read_jsonl(os.path.join(DATA, 'runs.jsonl'))
    for r in runs[-200:]:  # 외국어 Top3 제목 번역(캐시에 없는 것만 요청)
        for o in r['sources'].values():
            for t in (o.get('top3') or []):
                if 'title_ko' not in t:
                    t['title_ko'] = TR.to_ko(t['title'], 300)
    TR.save()
    steam = read_jsonl(os.path.join(DATA, 'steam.jsonl'))
    labels = {'labeled': len(labeled), 'files': files, 'usable': bool(nb and nb.usable),
                     'cv': (nb.cv if nb else None), 'model_kind': nb.kind if nb else 'lexicon',
                     'class_n': dict(nb.class_n) if nb else {},
                     'state': model_state(), 'pending': label_signature(labeled) != model_state().get('sig'),
                     'reasons': [{k: r.get(k) for k in ('label', 'reason_type', 'reason', 'src', 'pred')}
                                 for r in labeled if r.get('reason') or r.get('reason_type')],
                     'user_lex': S.USER_LEX, 'ai': ai_state(), 'ai_every': AI_RETEST_EVERY}
    if CLOUD:
        dashboard.build(BASE, CFG, SOURCES, steam, runs, labels, out_path=os.path.join(BASE, 'site', 'index.html'),
                        public=True)
        return
    dashboard.build(BASE, CFG, SOURCES, steam, runs, labels)
    publish_public(steam, runs, labels)


def publish_public(steam, runs, labels):
    """config 의 publish.repo 와 GitHub 토큰이 있으면 공개용 페이지(판단 이유 제외·noindex)를 GitHub Pages 로 올린다."""
    pub = CFG.get('publish') or {}
    if not pub.get('repo'):
        return
    import dashboard
    import publish
    if not publish.token():
        return
    out = os.path.join(DATA, 'public', 'index.html')
    try:
        dashboard.build(BASE, CFG, SOURCES, steam, runs, labels, out_path=out, public=True)
        res = publish.publish(out, pub['repo'], pub.get('branch', 'main'))
        if res == 'pushed':
            log(f"[publish] {publish.page_url(pub['repo'])} 갱신")
    except Exception as e:  # noqa
        log(f'[publish] 실패(수집에는 영향 없음): {e}')


CATCHUP_WINDOWS = 4  # 클라우드: 예약 실행이 늦어져도 최근 2시간 안의 빠진 30분 창은 채운다


def run_cloud():
    """GitHub Actions 에서 15분마다: Steam 기록 → 처리 안 된 최근 30분 창들을 오래된 것부터 수집 → 대시보드.
    GitHub 예약 실행은 몰리는 시간에 10~20분 늦게 시작하기도 해서 '현재 창 하나'가 아니라 빈 창을 찾아 채운다."""
    rc = run_steam()
    _, we_now = current_window()
    done = {r['slot_end'] for r in read_jsonl(os.path.join(DATA, 'runs.jsonl'))}
    todo = [we_now - WINDOW * k for k in range(CATCHUP_WINDOWS - 1, -1, -1)]
    for we in [w for w in todo if iso(w) not in done]:
        rc = max(rc, run_community(now=we + timedelta(seconds=30)) or 0)
    try:
        build_dashboard()
    except Exception as e:  # noqa
        log(f'[build] 실패: {e}')
    return rc


def keep_awake():
    """작업 스케줄러가 절전 중인 PC를 깨워 실행했을 때(WakeToRun), 5~6분 걸리는 수집 도중
    다시 잠들지 않도록 이 프로세스가 끝날 때까지만 시스템 절전을 보류한다(화면은 꺼진 채 유지)."""
    if os.name == 'nt':
        try:
            import ctypes
            ES_CONTINUOUS, ES_SYSTEM_REQUIRED = 0x80000000, 0x00000001
            ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS | ES_SYSTEM_REQUIRED)
        except Exception:  # noqa
            pass


MIN_BATTERY = int(CFG.get('min_battery_pct', 30))


def battery():
    """(전원 연결 여부, 잔량 %) — Windows GetSystemPowerStatus. 알 수 없으면 (True, None)."""
    if os.name != 'nt':
        return True, None
    try:
        import ctypes

        class SPS(ctypes.Structure):
            _fields_ = [('ACLineStatus', ctypes.c_ubyte), ('BatteryFlag', ctypes.c_ubyte),
                        ('BatteryLifePercent', ctypes.c_ubyte), ('SystemStatusFlag', ctypes.c_ubyte),
                        ('BatteryLifeTime', ctypes.c_ulong), ('BatteryFullLifeTime', ctypes.c_ulong)]
        s = SPS()
        if not ctypes.windll.kernel32.GetSystemPowerStatus(ctypes.byref(s)):
            return True, None
        pct = None if s.BatteryLifePercent == 255 else s.BatteryLifePercent
        return s.ACLineStatus != 0, pct
    except Exception:  # noqa
        return True, None


def main():
    # 배터리일 때도 절전 중 깨워 수집한다(2026-10-07 사용자 선택, 전원옵션 깨우기 타이머 DC=1).
    # 대신 잔량이 MIN_BATTERY% 미만이면 이 회차는 바로 건너뛴다 — 실패(exit 1)로 끝내면 스케줄러가
    # 재시도해 배터리를 더 쓰므로 정상(0)으로 끝낸다. 대시보드 '수집량' 카드에 빠진 창으로 보인다.
    ac, pct = battery()
    cmd0 = sys.argv[1] if len(sys.argv) > 1 else 'community'
    if cmd0 in ('steam', 'community') and not ac and pct is not None and pct < MIN_BATTERY:
        log(f'[battery] 배터리 {pct}% (< {MIN_BATTERY}%) · 전원 미연결 → {cmd0} 회차 건너뜀')
        return 0
    keep_awake()
    cmd = sys.argv[1] if len(sys.argv) > 1 else 'community'
    if cmd == 'steam':
        rc = run_steam()
        try:
            build_dashboard()
        except Exception as e:  # noqa
            log(f'[build] 실패: {e}')
        return rc
    if cmd == 'community':
        return run_community()
    if cmd == 'build':
        build_dashboard()
        return 0
    if cmd == 'learn':
        nb, labeled, files = build_model()
        print(f'라벨 {len(labeled)}건, 분포 {dict(nb.class_n) if nb else {}}, 모델: {nb.kind if nb else "lexicon"}')
        if nb and nb.cv:
            print(f"  5-fold: 사전만 {nb.cv['lexicon']}% · 로지스틱 최고 {nb.cv['logreg_best']}% → 채택 {nb.cv['best'] or '사전만'}")
        for f in files:
            print(f"  {f['file']}: {f['labeled']}건 라벨, 프로그램 일치 {f['agree']}")
        return 0
    if cmd == 'rescore':
        run_rescore()
        return 0
    if cmd == 'cloud':
        return run_cloud()
    if cmd == 'export-labels':
        out = sys.argv[2] if len(sys.argv) > 2 else LABELS_ENC
        print(f'라벨 {export_labels(out)}건 → {out}')
        return 0
    print(__doc__)
    return 2


if __name__ == '__main__':
    try:
        sys.exit(main())
    except Exception as e:  # noqa
        log(f'[fatal] {e}\n{traceback.format_exc()}')
        sys.exit(1)
